"""Receivables reports service: Ageing Summary, Ageing Detail, Client Statement, Collections, and WHT Reports."""
from __future__ import annotations

import calendar
import io
import logging
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.graphics.shapes import Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services import finance_pdf as finance_pdf_mod
from app.services.finance_pdf import NumberedCanvas, _wordmark_flowable, format_currency_amount

logger = logging.getLogger(__name__)



_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def spreadsheet_safe(value: Any) -> Any:
    """Neutralise spreadsheet formula injection (OWASP CSV injection) in exported text cells."""
    if isinstance(value, str) and value and value.startswith(_FORMULA_PREFIXES) and value.strip() not in ("-", "--"):
        return "'" + value
    return value


def _safe_append(ws: Any, row: List[Any]) -> None:
    ws.append([spreadsheet_safe(v) for v in row])


def _today_str() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def _as_int(value: Any, default: int = 0) -> int:
    try:
        if value is None:
            return default
        return int(value)
    except (TypeError, ValueError):
        return default


def _clean_mongo_doc(doc: Dict[str, Any]) -> Dict[str, Any]:
    cleaned = dict(doc)
    cleaned.pop("_id", None)
    return cleaned


def _calc_days_overdue(due_date_str: str, as_of_str: str) -> int:
    try:
        due = datetime.strptime(due_date_str, "%Y-%m-%d")
        as_of = datetime.strptime(as_of_str, "%Y-%m-%d")
        delta = (as_of - due).days
        return max(0, delta)
    except Exception:
        return 0


def _get_age_bucket(days: int) -> str:
    if days <= 0:
        return "current"
    elif days <= 30:
        return "1_30"
    elif days <= 60:
        return "31_60"
    elif days <= 90:
        return "61_90"
    else:
        return "90_plus"


async def get_ageing_summary_report(
    db,
    company_id: str,
    as_of_date: Optional[str] = None,
    client_name: Optional[str] = None,
    currency: Optional[str] = None,
) -> Dict[str, Any]:
    """Receivables Ageing Summary: Client rows x Age Buckets x Currencies."""
    as_of = as_of_date or _today_str()
    coll = db["invoices"]

    query: Dict[str, Any] = {
        "company_id": company_id,
        "issue_date": {"$lte": as_of},
        "status": {"$in": ["issued", "partially_paid", "overdue", "paid"]},
    }
    if client_name:
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}
    if currency and currency.upper() != "ALL":
        query["currency"] = currency.upper().strip()

    cursor = coll.find(query)
    invoices = await cursor.to_list(length=1000)

    # To calculate exact balance as of date:
    # balance_as_of = total - receipts posted on or before as_of date
    # (Receipts voided after as_of date still counted on as_of date!)
    clients_map: Dict[str, Dict[str, Any]] = {}
    currencies_summary: Dict[str, Dict[str, int]] = {}

    for inv in invoices:
        inv_id = inv["id"]
        c_name = inv.get("client_name", "Unknown Client")
        curr = inv.get("currency", "PKR")
        total = inv.get("total", 0)

        # Find receipts allocated to this invoice on or before as_of date
        rcpt_cursor = db["receipts"].find({
            "company_id": company_id,
            "date": {"$lte": as_of},
            "allocations.invoice_id": inv_id,
            "$or": [
                {"status": "posted"},
                {"status": "void", "voided_at": {"$gt": as_of}},
            ],
        })
        matched_receipts = await rcpt_cursor.to_list(length=50)

        paid_as_of = 0
        wht_as_of = 0
        for r in matched_receipts:
            for al in r.get("allocations", []):
                if al.get("invoice_id") == inv_id:
                    paid_as_of += al.get("amount", 0)
                    wht_as_of += al.get("wht", 0)

        bal_as_of = max(0, total - (paid_as_of + wht_as_of))
        if bal_as_of <= 0:
            continue

        days = _calc_days_overdue(inv.get("due_date", as_of), as_of)
        bucket = _get_age_bucket(days)

        client_key = f"{c_name}_{curr}"
        if client_key not in clients_map:
            clients_map[client_key] = {
                "client_name": c_name,
                "currency": curr,
                "current": 0,
                "1_30": 0,
                "31_60": 0,
                "61_90": 0,
                "90_plus": 0,
                "total": 0,
            }

        clients_map[client_key][bucket] += bal_as_of
        clients_map[client_key]["total"] += bal_as_of

        if curr not in currencies_summary:
            currencies_summary[curr] = {
                "current": 0,
                "1_30": 0,
                "31_60": 0,
                "61_90": 0,
                "90_plus": 0,
                "total": 0,
            }
        currencies_summary[curr][bucket] += bal_as_of
        currencies_summary[curr]["total"] += bal_as_of

    rows = sorted(clients_map.values(), key=lambda r: (r["currency"], -r["total"]))
    return {
        "as_of_date": as_of,
        "rows": rows,
        "currencies_summary": currencies_summary,
    }


async def get_ageing_detail_report(
    db,
    company_id: str,
    as_of_date: Optional[str] = None,
    client_name: Optional[str] = None,
    currency: Optional[str] = None,
) -> Dict[str, Any]:
    """Receivables Ageing Detail: Invoice rows grouped by Client."""
    as_of = as_of_date or _today_str()
    coll = db["invoices"]

    query: Dict[str, Any] = {
        "company_id": company_id,
        "issue_date": {"$lte": as_of},
        "status": {"$in": ["issued", "partially_paid", "overdue", "paid"]},
    }
    if client_name:
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}
    if currency and currency.upper() != "ALL":
        query["currency"] = currency.upper().strip()

    cursor = coll.find(query).sort("due_date", 1)
    invoices = await cursor.to_list(length=1000)

    groups: Dict[str, Dict[str, Any]] = {}
    for inv in invoices:
        inv_id = inv["id"]
        c_name = inv.get("client_name", "Unknown Client")
        curr = inv.get("currency", "PKR")
        total = inv.get("total", 0)

        rcpt_cursor = db["receipts"].find({
            "company_id": company_id,
            "date": {"$lte": as_of},
            "allocations.invoice_id": inv_id,
            "$or": [
                {"status": "posted"},
                {"status": "void", "voided_at": {"$gt": as_of}},
            ],
        })
        matched_receipts = await rcpt_cursor.to_list(length=50)

        paid_as_of = 0
        wht_as_of = 0
        for r in matched_receipts:
            for al in r.get("allocations", []):
                if al.get("invoice_id") == inv_id:
                    paid_as_of += al.get("amount", 0)
                    wht_as_of += al.get("wht", 0)

        bal_as_of = max(0, total - (paid_as_of + wht_as_of))
        if bal_as_of <= 0:
            continue

        days = _calc_days_overdue(inv.get("due_date", as_of), as_of)

        group_key = f"{c_name}_{curr}"
        if group_key not in groups:
            groups[group_key] = {
                "client_name": c_name,
                "currency": curr,
                "invoices": [],
                "subtotal_total": 0,
                "subtotal_paid": 0,
                "subtotal_wht": 0,
                "subtotal_balance": 0,
            }

        inv_row = {
            "id": inv["id"],
            "number": inv.get("number"),
            "issue_date": inv.get("issue_date"),
            "due_date": inv.get("due_date"),
            "days_overdue": days,
            "total": total,
            "amount_paid": paid_as_of,
            "wht_settled": wht_as_of,
            "balance_due": bal_as_of,
            "status": "overdue" if days > 0 else "current",
        }
        groups[group_key]["invoices"].append(inv_row)
        groups[group_key]["subtotal_total"] += total
        groups[group_key]["subtotal_paid"] += paid_as_of
        groups[group_key]["subtotal_wht"] += wht_as_of
        groups[group_key]["subtotal_balance"] += bal_as_of

    return {"as_of_date": as_of, "groups": list(groups.values())}


async def get_client_statement_report(
    db,
    company_id: str,
    client_name: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    currency: Optional[str] = "PKR",
) -> Dict[str, Any]:
    """Client Statement: Opening balance + Chronological Invoices & Receipts with running balance."""
    start_date = date_from or "2020-01-01"
    end_date = date_to or _today_str()
    curr = (currency or "PKR").upper().strip()

    # 1. Calculate Opening Balance before start_date
    prior_inv_cursor = db["invoices"].find({
        "company_id": company_id,
        "client_name": {"$regex": f"^{re.escape(client_name.strip())}$", "$options": "i"},
        "currency": curr,
        "issue_date": {"$lt": start_date},
        "status": {"$in": ["issued", "partially_paid", "overdue", "paid"]},
    })
    prior_invoices = await prior_inv_cursor.to_list(length=500)
    prior_inv_total = sum(i.get("total", 0) for i in prior_invoices)

    prior_rcpt_cursor = db["receipts"].find({
        "company_id": company_id,
        "client_name": {"$regex": f"^{re.escape(client_name.strip())}$", "$options": "i"},
        "currency": curr,
        "date": {"$lt": start_date},
        "status": "posted",
    })
    prior_receipts = await prior_rcpt_cursor.to_list(length=500)
    prior_settled = sum(r.get("settles_total", 0) for r in prior_receipts)
    opening_balance = max(0, prior_inv_total - prior_settled)

    # 2. Collect Invoices and Receipts in Date Range
    period_invoices_cursor = db["invoices"].find({
        "company_id": company_id,
        "client_name": {"$regex": f"^{re.escape(client_name.strip())}$", "$options": "i"},
        "currency": curr,
        "issue_date": {"$gte": start_date, "$lte": end_date},
        "status": {"$in": ["issued", "partially_paid", "overdue", "paid"]},
    })
    period_invoices = await period_invoices_cursor.to_list(length=500)

    period_receipts_cursor = db["receipts"].find({
        "company_id": company_id,
        "client_name": {"$regex": f"^{re.escape(client_name.strip())}$", "$options": "i"},
        "currency": curr,
        "date": {"$gte": start_date, "$lte": end_date},
        "status": "posted",
    })
    period_receipts = await period_receipts_cursor.to_list(length=500)

    # 3. Interleave and Sort Chronologically
    transactions = []
    for inv in period_invoices:
        transactions.append({
            "type": "invoice",
            "date": inv.get("issue_date"),
            "number": inv.get("number"),
            "description": f"Invoice {inv.get('number')} (Due: {inv.get('due_date')})",
            "debit": inv.get("total", 0),
            "credit": 0,
        })
    for rcpt in period_receipts:
        memo = f"Receipt {rcpt.get('number')} via {rcpt.get('method', '').upper()}"
        if rcpt.get("reference_number"):
            memo += f" (Ref: {rcpt.get('reference_number')})"
        transactions.append({
            "type": "receipt",
            "date": rcpt.get("date"),
            "number": rcpt.get("number"),
            "description": memo,
            "debit": 0,
            "credit": rcpt.get("settles_total", 0),
        })

    transactions.sort(key=lambda t: t["date"])

    # Compute Running Balance
    running_balance = opening_balance
    for t in transactions:
        running_balance = running_balance + t["debit"] - t["credit"]
        t["running_balance"] = running_balance

    return {
        "client_name": client_name,
        "currency": curr,
        "date_from": start_date,
        "date_to": end_date,
        "opening_balance": opening_balance,
        "closing_balance": running_balance,
        "transactions": transactions,
    }


async def get_collections_report(
    db,
    company_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    method: Optional[str] = None,
    bank: Optional[str] = None,
    client_name: Optional[str] = None,
) -> Dict[str, Any]:
    """Collections breakdown by method and bank account."""
    start_date = date_from or "2020-01-01"
    end_date = date_to or _today_str()

    query: Dict[str, Any] = {
        "company_id": company_id,
        "date": {"$gte": start_date, "$lte": end_date},
        "status": {"$in": ["posted", "pending"]},
    }
    if method and method.lower() != "all":
        query["method"] = method.lower().strip()
    if bank and bank.lower() != "all":
        query["bank"] = bank.lower().strip()
    if client_name and client_name.strip():
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    cursor = db["receipts"].find(query).sort("date", -1)
    raw_receipts = await cursor.to_list(length=1000)
    receipts = [_clean_mongo_doc(r) for r in raw_receipts]

    summary = {
        "cash_pkr": 0,
        "abl_pkr": 0,
        "meezan_pkr": 0,
        "total_wht_pkr": 0,
        "total_collected_pkr": 0,
        "pending_pkr": 0,
        "on_hold_pkr": 0,
    }

    for r in receipts:
        pkr = _as_int(r.get("pkr_amount"))
        m = r.get("method")
        b = r.get("bank")
        st = r.get("status")

        if st == "pending":
            summary["pending_pkr"] += pkr
            continue
        if st == "void":
            summary["on_hold_pkr"] += pkr
            continue

        summary["total_collected_pkr"] += pkr
        if m == "cash":
            summary["cash_pkr"] += pkr
        elif b == "abl":
            summary["abl_pkr"] += pkr
        elif b == "meezan":
            summary["meezan_pkr"] += pkr

        wht = _as_int(r.get("wht_amount"))
        if r.get("currency") == "PKR":
            summary["total_wht_pkr"] += wht
        elif r.get("exchange_rate"):
            try:
                summary["total_wht_pkr"] += int(round(wht * float(r.get("exchange_rate") or 0)))
            except (TypeError, ValueError):
                pass

    return {
        "date_from": start_date,
        "date_to": end_date,
        "receipts": receipts,
        "summary": summary,
    }


async def get_receivables_outstanding_report(
    db,
    company_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    client_name: Optional[str] = None,
) -> Dict[str, Any]:
    """Outstanding balances by client for invoices issued in the selected range."""
    start_date = date_from or "2020-01-01"
    end_date = date_to or _today_str()
    query: Dict[str, Any] = {
        "company_id": company_id,
        "issue_date": {"$gte": start_date, "$lte": end_date},
        "status": {"$in": ["issued", "partially_paid", "overdue"]},
        "balance_due": {"$gt": 0},
    }
    if client_name and client_name.strip():
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    invoices = await db["invoices"].find(query).to_list(length=1000)
    grouped: Dict[str, Dict[str, Any]] = {}
    for inv in invoices:
        curr = inv.get("currency") or "PKR"
        name = inv.get("client_name") or "Unknown Client"
        key = f"{name}::{curr}"
        row = grouped.get(key)
        if row:
            row["total"] += _as_int(inv.get("balance_due"))
            row["invoice_count"] += 1
        else:
            grouped[key] = {
                "client_name": name,
                "currency": curr,
                "total": _as_int(inv.get("balance_due")),
                "invoice_count": 1,
            }
    rows = sorted(grouped.values(), key=lambda r: r["client_name"].lower())
    totals: Dict[str, int] = {}
    for row in rows:
        totals[row["currency"]] = totals.get(row["currency"], 0) + row["total"]
    return {
        "date_from": start_date,
        "date_to": end_date,
        "rows": rows,
        "totals": totals,
    }


async def get_summary_accounts_report(
    db,
    company_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    client_name: Optional[str] = None,
) -> Dict[str, Any]:
    """Invoice, receipts, and remaining balance by issued invoice in the selected range."""
    start_date = date_from or "2020-01-01"
    end_date = date_to or _today_str()
    query: Dict[str, Any] = {
        "company_id": company_id,
        "issue_date": {"$gte": start_date, "$lte": end_date},
        "status": {"$nin": ["void", "deleted", "draft"]},
    }
    if client_name and client_name.strip():
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    invoices = await db["invoices"].find(query).sort("issue_date", 1).to_list(length=1000)
    rows = []
    totals_by_currency: Dict[str, Dict[str, int]] = {}
    for inv in invoices:
        curr = inv.get("currency") or "PKR"
        invoice_amt = _as_int(inv.get("total"))
        received = _as_int(inv.get("amount_paid")) + _as_int(inv.get("wht_settled"))
        due = _as_int(inv.get("balance_due"))
        rows.append({
            "issue_date": inv.get("issue_date") or "",
            "number": inv.get("number") or "",
            "client_name": inv.get("client_name") or "Unknown Client",
            "currency": curr,
            "invoice_amount": invoice_amt,
            "received": received,
            "balance_due": due,
        })
        bucket = totals_by_currency.setdefault(curr, {"invoice": 0, "received": 0, "due": 0})
        bucket["invoice"] += invoice_amt
        bucket["received"] += received
        bucket["due"] += due
    return {
        "date_from": start_date,
        "date_to": end_date,
        "rows": rows,
        "totals": totals_by_currency,
    }


async def get_wht_report(
    db,
    company_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    client_name: Optional[str] = None,
) -> Dict[str, Any]:
    """WHT withholding tracker with certificate verification status."""
    start_date = date_from or "2020-01-01"
    end_date = date_to or _today_str()

    query: Dict[str, Any] = {
        "company_id": company_id,
        "date": {"$gte": start_date, "$lte": end_date},
        "wht_amount": {"$gt": 0},
        "status": {"$in": ["posted", "pending"]},
    }
    if client_name:
        query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    cursor = db["receipts"].find(query).sort("date", -1)
    records = await cursor.to_list(length=500)

    total_wht_pkr = 0
    total_with_cert = 0
    total_pending_cert = 0

    items = []
    for r in records:
        wht_amt = r.get("wht_amount", 0)
        has_cert = bool(r.get("wht_certificate_file_id"))
        if has_cert:
            total_with_cert += 1
        else:
            total_pending_cert += 1

        if r.get("currency") == "PKR":
            wht_pkr = wht_amt
        else:
            rate = r.get("exchange_rate") or 1.0
            wht_pkr = int(round(wht_amt * rate))
        total_wht_pkr += wht_pkr

        items.append({
            "receipt_id": r["id"],
            "receipt_number": r.get("number"),
            "date": r.get("date"),
            "client_name": r.get("client_name"),
            "currency": r.get("currency"),
            "wht_amount": wht_amt,
            "wht_pkr": wht_pkr,
            "has_certificate": has_cert,
            "certificate_file_id": r.get("wht_certificate_file_id"),
        })

    return {
        "date_from": start_date,
        "date_to": end_date,
        "total_wht_pkr": total_wht_pkr,
        "total_with_cert": total_with_cert,
        "total_pending_cert": total_pending_cert,
        "items": items,
    }


# Posted collections only. 50,000 PKR (minor units) or 5% — whichever is larger.
_TREND_DEAD_BAND_MINOR = 5_000_000
_TREND_DEAD_BAND_PCT = 0.05
_MONTH_RE = re.compile(r"^\d{4}-\d{2}$")


def _month_key(date_str: Optional[str]) -> str:
    raw = (date_str or "").strip()
    return raw[:7] if len(raw) >= 7 else ""


def _parse_month(month: str) -> Tuple[int, int]:
    year_s, month_s = month.split("-", 1)
    return int(year_s), int(month_s)


def _prev_month(month: str) -> str:
    year, mon = _parse_month(month)
    if mon == 1:
        return f"{year - 1:04d}-12"
    return f"{year:04d}-{mon - 1:02d}"


def _month_label(month: str) -> str:
    try:
        return datetime.strptime(f"{month}-01", "%Y-%m-%d").strftime("%b %Y")
    except ValueError:
        return month


def _iter_months(date_from: str, date_to: str) -> List[str]:
    start = _month_key(date_from)
    end = _month_key(date_to)
    if not start or not end or start > end:
        return []
    year, mon = _parse_month(start)
    end_year, end_mon = _parse_month(end)
    months: List[str] = []
    while (year, mon) <= (end_year, end_mon):
        months.append(f"{year:04d}-{mon:02d}")
        if mon == 12:
            year += 1
            mon = 1
        else:
            mon += 1
    return months


def _days_in_month(month: str) -> int:
    year, mon = _parse_month(month)
    return calendar.monthrange(year, mon)[1]


PERIODS = ("monthly", "quarterly", "semiannual", "yearly")
_PERIOD_MONTHS = {"monthly": 1, "quarterly": 3, "semiannual": 6, "yearly": 12}
_PERIOD_KEY_RE = re.compile(r"^\d{4}(?:-(?:0[1-9]|1[0-2]|Q[1-4]|H[12]))?$")


def normalize_period(period: Optional[str]) -> str:
    p = (period or "monthly").strip().lower().replace("-", "").replace("_", "")
    if p in ("semiannually", "halfyearly", "halfyear", "6m", "biannual"):
        p = "semiannual"
    if p in ("quarter", "3m"):
        p = "quarterly"
    if p in ("year", "annual", "annually", "12m"):
        p = "yearly"
    if p in ("month", "1m"):
        p = "monthly"
    if p not in PERIODS:
        raise ValueError(f"period must be one of {', '.join(PERIODS)}")
    return p


def is_period_key(key: Optional[str]) -> bool:
    return bool(key and _PERIOD_KEY_RE.match(key.strip()))


def _period_of(date_str: Optional[str], period: str = "monthly") -> str:
    month = _month_key(date_str)
    if not month:
        return ""
    try:
        year, mon = _parse_month(month)
    except ValueError:
        return ""
    if period == "quarterly":
        return f"{year:04d}-Q{(mon - 1) // 3 + 1}"
    if period == "semiannual":
        return f"{year:04d}-H{1 if mon <= 6 else 2}"
    if period == "yearly":
        return f"{year:04d}"
    return month


def _period_months(key: str) -> List[str]:
    """Calendar months (YYYY-MM) covered by a period key."""
    year = int(key[:4])
    if len(key) == 4:
        start, n = 1, 12
    elif key[5] == "Q":
        start, n = (int(key[6]) - 1) * 3 + 1, 3
    elif key[5] == "H":
        start, n = 1 if key[6] == "1" else 7, 6
    else:
        start, n = int(key[5:7]), 1
    return [f"{year:04d}-{m:02d}" for m in range(start, start + n)]


def _period_first_day(key: str) -> str:
    return f"{_period_months(key)[0]}-01"


def _period_last_day(key: str) -> str:
    last = _period_months(key)[-1]
    return f"{last}-{_days_in_month(last):02d}"


def _days_in_period(key: str) -> int:
    return sum(_days_in_month(m) for m in _period_months(key))


def period_start(date_str: str, period: str) -> str:
    """First day of the period containing date_str (the range snaps outward to whole periods)."""
    key = _period_of(date_str, period)
    return _period_first_day(key) if key else date_str


def _prev_period(key: str) -> str:
    if len(key) == 7 and key[4] == "-" and key[5].isdigit():
        return _prev_month(key)
    first = _period_months(key)[0]
    prev_month = _prev_month(first)
    if len(key) == 4:
        return prev_month[:4]
    return _period_of(f"{prev_month}-01", "quarterly" if key[5] == "Q" else "semiannual")


def _period_key_label(key: str) -> str:
    if len(key) == 4:
        return key
    if key[5] in ("Q", "H"):
        return f"{key[5:]} {key[:4]}"
    return _month_label(key)


def _iter_periods(date_from: str, date_to: str, period: str = "monthly") -> List[str]:
    keys: List[str] = []
    for month in _iter_months(date_from, date_to):
        key = _period_of(f"{month}-01", period)
        if not keys or keys[-1] != key:
            keys.append(key)
    return keys


def run_rate_pct(actual: int, run_rate: Optional[int]) -> Optional[float]:
    """% the selected period is above (+) or below (-) its pace-adjusted 3-period run rate.

    None when there is no run rate yet or it is zero (avoids divide-by-zero).
    """
    if run_rate is None or run_rate <= 0:
        return None
    return round(((actual - run_rate) / run_rate) * 1000) / 10.0


def _receipt_wht_pkr(receipt: Dict[str, Any]) -> int:
    wht = _as_int(receipt.get("wht_amount"))
    if receipt.get("currency") == "PKR":
        return wht
    try:
        return int(round(wht * float(receipt.get("exchange_rate") or 0)))
    except (TypeError, ValueError):
        return 0


def _receipt_channel(receipt: Dict[str, Any]) -> str:
    if receipt.get("method") == "cash":
        return "cash"
    bank = (receipt.get("bank") or "").lower().strip()
    if bank in ("abl", "meezan"):
        return bank
    return "other"


def _empty_month_bucket(month: str, period: str = "monthly") -> Dict[str, Any]:
    return {
        "month": month,
        "label": _period_key_label(month) if month and month != "0000-00" else _month_label(month),
        "collected_pkr": 0,
        "billed_pkr": 0,
        "pending_pkr": 0,
        "wht_pkr": 0,
        "avg_3m_pkr": 0,
        "is_partial": False,
        "collected_new_invoices_pkr": 0,
        "collected_prior_invoices_pkr": 0,
        "collected_unallocated_pkr": 0,
        "channels": {"cash": 0, "abl": 0, "meezan": 0, "other": 0},
    }


def _classify_direction(delta: int, prior: int) -> str:
    if prior <= 0 and delta == 0:
        return "flat"
    if prior <= 0 and delta > 0:
        return "improving"
    band = max(_TREND_DEAD_BAND_MINOR, int(round(abs(prior) * _TREND_DEAD_BAND_PCT)))
    if abs(delta) <= band:
        return "flat"
    return "improving" if delta > 0 else "down"


def _run_rate_status(actual: int, run_rate: Optional[int]) -> Optional[str]:
    if run_rate is None or run_rate <= 0:
        return None
    band = max(_TREND_DEAD_BAND_MINOR, int(round(run_rate * _TREND_DEAD_BAND_PCT)))
    if abs(actual - run_rate) <= band:
        return "in_line"
    return "above" if actual > run_rate else "below"


def build_monthly_collections_analysis(
    *,
    date_from: str,
    date_to: str,
    today: str,
    receipts: List[Dict[str, Any]],
    invoices: List[Dict[str, Any]],
    focus_month: Optional[str] = None,
    period: str = "monthly",
) -> Dict[str, Any]:
    """Compare posted cash-in per period (month, quarter, half-year or year).

    Pending cheques are excluded from collected. Keys keep their historical
    `month` names for compatibility; with a non-monthly period they hold period
    keys such as 2026-Q4, 2026-H2 or 2026.
    """
    period = normalize_period(period)
    date_from = period_start(date_from, period) if date_from else date_from
    months = _iter_periods(date_from, date_to, period)
    buckets = {month: _empty_month_bucket(month, period) for month in months}
    client_by_month: Dict[str, Dict[str, int]] = {month: {} for month in months}

    for inv in invoices:
        month = _period_of(inv.get("issue_date"), period)
        bucket = buckets.get(month)
        if not bucket:
            continue
        if (inv.get("currency") or "PKR") == "PKR":
            bucket["billed_pkr"] += _as_int(inv.get("total"))

    for receipt in receipts:
        month = _period_of(receipt.get("date"), period)
        bucket = buckets.get(month)
        if not bucket:
            continue
        status = receipt.get("status")
        pkr = _as_int(receipt.get("pkr_amount"))
        if status == "pending":
            bucket["pending_pkr"] += pkr
            continue
        if status != "posted":
            continue

        bucket["collected_pkr"] += pkr
        bucket["wht_pkr"] += _receipt_wht_pkr(receipt)
        channel = _receipt_channel(receipt)
        bucket["channels"][channel] = bucket["channels"].get(channel, 0) + pkr
        client_name = (receipt.get("client_name") or "Unknown client").strip() or "Unknown client"
        client_by_month[month][client_name] = client_by_month[month].get(client_name, 0) + pkr

        allocs = receipt.get("allocations") or []
        allocated = 0
        for alloc in allocs:
            amt = _as_int(alloc.get("amount"))
            if (receipt.get("currency") or "PKR") != "PKR":
                try:
                    amt = int(round(amt * float(receipt.get("exchange_rate") or 0)))
                except (TypeError, ValueError):
                    amt = 0
            allocated += amt
            inv_month = _period_of(alloc.get("invoice_date"), period)
            if inv_month and inv_month == month:
                bucket["collected_new_invoices_pkr"] += amt
            else:
                bucket["collected_prior_invoices_pkr"] += amt
        leftover = pkr - allocated
        if leftover > 0:
            bucket["collected_unallocated_pkr"] += leftover

    cutoff = min(date_to, today) if today else date_to
    cutoff_month = _period_of(cutoff, period)
    series: List[Dict[str, Any]] = []
    for idx, month in enumerate(months):
        bucket = buckets[month]
        last_day = _period_last_day(month)
        bucket["is_partial"] = month == cutoff_month and cutoff < last_day
        window = [buckets[m]["collected_pkr"] for m in months[max(0, idx - 2) : idx + 1]]
        bucket["avg_3m_pkr"] = int(round(sum(window) / len(window))) if window else 0
        series.append(bucket)

    requested = (focus_month or "").strip()
    if requested and requested in buckets:
        selected_month = requested
    else:
        end_key = _period_of(date_to, period)
        selected_month = end_key if end_key in buckets else (months[-1] if months else "")

    selected = buckets.get(selected_month) or _empty_month_bucket(selected_month or "0000-00", period)
    prior_key = _prev_period(selected_month) if selected_month else ""
    prior = buckets.get(prior_key)
    has_prior = prior is not None
    prior_collected = prior["collected_pkr"] if prior else 0
    delta = selected["collected_pkr"] - prior_collected
    pct_change = None
    if has_prior and prior_collected > 0:
        pct_change = round((delta / prior_collected) * 1000) / 10.0
    elif has_prior and selected["collected_pkr"] > 0:
        pct_change = 100.0

    direction = "no_prior" if not has_prior else _classify_direction(delta, prior_collected)

    complete_before = [
        buckets[m]["collected_pkr"]
        for m in months
        if m < selected_month and not buckets[m]["is_partial"]
    ]
    run_rate_source = complete_before[-3:]
    run_rate_full = (
        int(round(sum(run_rate_source) / len(run_rate_source))) if run_rate_source else None
    )
    dim = _days_in_period(selected_month) if selected_month else 30
    if selected["is_partial"]:
        try:
            elapsed = (
                datetime.strptime(cutoff[:10], "%Y-%m-%d")
                - datetime.strptime(_period_first_day(selected_month), "%Y-%m-%d")
            ).days + 1
        except (TypeError, ValueError):
            elapsed = dim
        elapsed = min(max(elapsed, 1), dim)
        run_rate_compare = int(round(run_rate_full * elapsed / dim)) if run_rate_full is not None else None
    else:
        elapsed = dim
        run_rate_compare = run_rate_full

    selected_clients = client_by_month.get(selected_month) or {}
    prior_clients = client_by_month.get(prior_key) or {}
    names = set(selected_clients) | set(prior_clients)
    movers = []
    for name in names:
        now_amt = selected_clients.get(name, 0)
        prev_amt = prior_clients.get(name, 0)
        client_delta = now_amt - prev_amt
        if client_delta == 0:
            continue
        movers.append({
            "client_name": name,
            "this_month_pkr": now_amt,
            "prior_month_pkr": prev_amt,
            "delta_pkr": client_delta,
        })
    movers.sort(key=lambda row: abs(row["delta_pkr"]), reverse=True)
    improved = [row for row in movers if row["delta_pkr"] > 0][:8]
    declined = [row for row in movers if row["delta_pkr"] < 0][:8]

    top_client_pkr = max(selected_clients.values()) if selected_clients else 0
    concentration_pct = (
        round((top_client_pkr / selected["collected_pkr"]) * 1000) / 10.0
        if selected["collected_pkr"] > 0
        else 0.0
    )

    return {
        "period": period,
        "date_from": date_from,
        "date_to": date_to,
        "headline_metric": "posted_collections_pkr",
        "headline_rule": (
            "Cash collected is posted receipts in PKR. Pending cheques are shown separately "
            "and are not treated as money in. Invoice totals are billed, not collected."
        ),
        "selected": {
            "month": selected_month,
            "label": selected["label"],
            "is_partial": selected["is_partial"],
            "days_elapsed": elapsed if selected_month else 0,
            "days_in_month": dim if selected_month else 0,
            "collected_pkr": selected["collected_pkr"],
            "prior_month": prior_key if has_prior else None,
            "prior_label": _period_key_label(prior_key) if has_prior else None,
            "prior_collected_pkr": prior_collected if has_prior else None,
            "delta_pkr": delta if has_prior else None,
            "pct_change": pct_change,
            "direction": direction,
            "billed_pkr": selected["billed_pkr"],
            "pending_pkr": selected["pending_pkr"],
            "wht_pkr": selected["wht_pkr"],
            "run_rate_3m_pkr": run_rate_full,
            "run_rate_compare_pkr": run_rate_compare,
            "run_rate_status": _run_rate_status(selected["collected_pkr"], run_rate_compare),
            "run_rate_pct": run_rate_pct(selected["collected_pkr"], run_rate_compare),
            "channels": selected["channels"],
            "collected_new_invoices_pkr": selected["collected_new_invoices_pkr"],
            "collected_prior_invoices_pkr": selected["collected_prior_invoices_pkr"],
            "collected_unallocated_pkr": selected["collected_unallocated_pkr"],
            "top_client_share_pct": concentration_pct,
        },
        "months": series,
        "movers_up": improved,
        "movers_down": declined,
    }


def _pkr_minor(rupees: int) -> int:
    return int(rupees) * 100


async def get_monthly_collections_analysis(
    db,
    company_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    client_name: Optional[str] = None,
    method: Optional[str] = None,
    bank: Optional[str] = None,
    focus_month: Optional[str] = None,
    period: str = "monthly",
) -> Dict[str, Any]:
    """Posted collections per period, period-over-period direction, and client movers."""
    period = normalize_period(period)
    end_date = date_to or _today_str()
    # Snap the start to the first day of its period so the first bucket is whole.
    start_date = period_start(date_from or "2020-01-01", period)
    receipt_query: Dict[str, Any] = {
        "company_id": company_id,
        "date": {"$gte": start_date, "$lte": end_date},
        "status": {"$in": ["posted", "pending"]},
    }
    if method and method.lower() != "all":
        receipt_query["method"] = method.lower().strip()
    if bank and bank.lower() != "all":
        receipt_query["bank"] = bank.lower().strip()
    if client_name and client_name.strip():
        receipt_query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    invoice_query: Dict[str, Any] = {
        "company_id": company_id,
        "issue_date": {"$gte": start_date, "$lte": end_date},
        "status": {"$nin": ["void", "deleted", "draft"]},
    }
    if client_name and client_name.strip():
        invoice_query["client_name"] = {"$regex": re.escape(client_name.strip()), "$options": "i"}

    receipt_projection = {
        "date": 1,
        "status": 1,
        "pkr_amount": 1,
        "wht_amount": 1,
        "currency": 1,
        "exchange_rate": 1,
        "client_name": 1,
        "method": 1,
        "bank": 1,
        "allocations": 1,
    }
    invoice_projection = {"issue_date": 1, "currency": 1, "total": 1, "status": 1}
    receipts = await db["receipts"].find(receipt_query, receipt_projection).to_list(length=20_000)
    invoices = await db["invoices"].find(invoice_query, invoice_projection).to_list(length=20_000)
    return build_monthly_collections_analysis(
        date_from=start_date,
        date_to=end_date,
        today=_today_str(),
        receipts=receipts,
        invoices=invoices,
        focus_month=focus_month,
        period=period,
    )


def _rupees_to_lakh(minor: Any) -> float:
    return _as_int(minor) / 100.0 / 100_000.0


def build_monthly_trend_chart(months: List[Dict[str, Any]], width: float, height: float) -> Optional[Drawing]:
    """Bar chart of collected vs billed (PKR lakh) with a 3-month average line."""
    if not months:
        return None
    collected = [_rupees_to_lakh(row.get("collected_pkr")) for row in months]
    billed = [_rupees_to_lakh(row.get("billed_pkr")) for row in months]
    avg = [_rupees_to_lakh(row.get("avg_3m_pkr")) for row in months]
    peak = max(collected + billed + avg + [1.0])
    y_max = max(2.0, round(peak * 1.15 * 2) / 2)
    drawing = Drawing(width, height)
    left, right, top, bottom = 36, 12, 18, 28
    plot_w = max(width - left - right, 40)
    plot_h = max(height - top - bottom, 40)
    ink = colors.HexColor("#18181B")
    muted = colors.HexColor("#71717A")
    rule = colors.HexColor("#E4E4E7")
    collected_color = colors.HexColor("#3F3F46")
    billed_color = colors.HexColor("#D4D4D8")
    drawing.add(Rect(0, 0, width, height, fillColor=colors.white, strokeColor=None))
    ticks = 4
    for i in range(ticks + 1):
        value = y_max * i / ticks
        y = bottom + plot_h * (value / y_max)
        drawing.add(Line(left, y, left + plot_w, y, strokeColor=rule, strokeWidth=0.4))
        drawing.add(String(0, y - 3, f"{value:.1f}", fontName="Helvetica", fontSize=6.5, fillColor=muted))
    drawing.add(String(0, height - 12, "PKR lakh", fontName="Helvetica", fontSize=6.5, fillColor=muted))
    n = len(months)
    slot = plot_w / n
    avg_points = []
    for idx, row in enumerate(months):
        x0 = left + idx * slot
        center = x0 + slot / 2
        billed_h = plot_h * (billed[idx] / y_max)
        collected_h = plot_h * (collected[idx] / y_max)
        bar_w = min(10, slot * 0.28)
        drawing.add(Rect(center - bar_w - 1, bottom, bar_w, billed_h, fillColor=billed_color, strokeColor=None))
        drawing.add(Rect(center + 1, bottom, bar_w, collected_h, fillColor=collected_color, strokeColor=None))
        avg_points.extend([center, bottom + plot_h * (avg[idx] / y_max)])
        label = (row.get("label") or "")[:3]
        if row.get("is_partial"):
            label += "*"
        drawing.add(String(center - 8, 8, label, fontName="Helvetica", fontSize=6.5, fillColor=muted))
    if len(avg_points) >= 4:
        drawing.add(PolyLine(avg_points, strokeColor=ink, strokeWidth=1.1, strokeDashArray=[3, 2]))
    legend_y = height - 8
    drawing.add(Rect(left + 70, legend_y - 3, 8, 8, fillColor=collected_color, strokeColor=None))
    drawing.add(String(left + 82, legend_y - 2, "Collected", fontName="Helvetica", fontSize=7, fillColor=muted))
    drawing.add(Rect(left + 140, legend_y - 3, 8, 8, fillColor=billed_color, strokeColor=None))
    drawing.add(String(left + 152, legend_y - 2, "Billed", fontName="Helvetica", fontSize=7, fillColor=muted))
    drawing.add(Line(left + 190, legend_y + 1, left + 208, legend_y + 1, strokeColor=ink, strokeWidth=1.1, strokeDashArray=[3, 2]))
    drawing.add(String(left + 212, legend_y - 2, "3-period avg", fontName="Helvetica", fontSize=7, fillColor=muted))
    return drawing


def export_monthly_analysis_pdf(report: Dict[str, Any], settings: Dict[str, Any]) -> bytes:
    selected = report.get("selected") or {}
    _unit = {"monthly": "month", "quarterly": "quarter", "semiannual": "half-year", "yearly": "year"}.get(report.get("period") or "monthly", "month")
    month_rows = []
    for row in report.get("months") or []:
        month_rows.append([
            f"{row.get('label') or ''}{' *' if row.get('is_partial') else ''}",
            format_currency_amount(_as_int(row.get("collected_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("billed_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("pending_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("wht_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("avg_3m_pkr")), "PKR"),
        ])
    movers = []
    for row in report.get("movers_up") or []:
        movers.append([
            row.get("client_name") or "",
            "Up",
            format_currency_amount(_as_int(row.get("prior_month_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("this_month_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("delta_pkr")), "PKR"),
        ])
    for row in report.get("movers_down") or []:
        movers.append([
            row.get("client_name") or "",
            "Down",
            format_currency_amount(_as_int(row.get("prior_month_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("this_month_pkr")), "PKR"),
            format_currency_amount(_as_int(row.get("delta_pkr")), "PKR"),
        ])
    direction = (selected.get("direction") or "flat").replace("_", " ")
    kpis = [
        ("Cash collected", format_currency_amount(_as_int(selected.get("collected_pkr")), "PKR")),
        (f"vs prior {_unit}", format_currency_amount(_as_int(selected.get("delta_pkr")), "PKR")),
        ("Direction", direction.title()),
        ("Pending cheques", format_currency_amount(_as_int(selected.get("pending_pkr")), "PKR")),
    ]
    rr_pct = selected.get("run_rate_pct")
    if rr_pct is not None:
        kpis.append(("vs run rate", f"{rr_pct:+.1f}%"))
    if movers:
        month_rows.append(["", "", "", "", "", ""])
        month_rows.append(["Client movers", "Change", f"Prior {_unit}", f"This {_unit}", "", ""])
        for mover in movers:
            month_rows.append([
                f"{mover[0]} ({mover[1]})",
                mover[4],
                mover[2],
                mover[3],
                "",
                "",
            ])
    extra: List[Any] = []
    chart = build_monthly_trend_chart(list(report.get("months") or []), 720, 210)
    if chart is not None:
        extra.append(chart)
        extra.append(Spacer(1, 12))
    return _standard_report_pdf(
        title="Collections report" if _unit != "month" else "Monthly collections report",
        period_label=_period_label(report.get("date_from") or "", report.get("date_to") or ""),
        settings=settings,
        headers=[_unit.capitalize(), "Collected", "Billed (PKR)", "Pending", "WHT", f"3-{_unit} avg"],
        col_weights=[0.18, 0.17, 0.17, 0.16, 0.16, 0.16],
        rows=month_rows,
        right_align_from=1,
        kpi_items=kpis,
        landscape_page=True,
        extra_flowables=extra or None,
    )


def export_ageing_summary_excel(report: Dict[str, Any]) -> bytes:
    """Generate professional Excel workbook for Ageing Summary report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ageing Summary"

    # Styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2B2D42", end_color="2B2D42", fill_type="solid")
    sub_fill = PatternFill(start_color="E9E5F5", end_color="E9E5F5", fill_type="solid")
    bold_font = Font(name="Calibri", size=10, bold=True)
    num_fmt = "#,##0"

    _safe_append(ws, [f"Receivables Ageing Summary (As of {report.get('as_of_date')})"])
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="2B2D42")
    _safe_append(ws, [])

    headers = ["Client Name", "Currency", "Current", "1-30 Days", "31-60 Days", "61-90 Days", "90+ Days", "Total"]
    _safe_append(ws, headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col_idx > 2 else "left")

    row_num = 4
    for r in report.get("rows", []):
        _safe_append(ws, [
            r["client_name"],
            r["currency"],
            r["current"] / 100.0,
            r["1_30"] / 100.0,
            r["31_60"] / 100.0,
            r["61_90"] / 100.0,
            r["90_plus"] / 100.0,
            r["total"] / 100.0,
        ])
        for c_idx in range(3, 9):
            ws.cell(row=row_num, column=c_idx).number_format = num_fmt
        row_num += 1

    # Totals per currency
    _safe_append(ws, [])
    row_num += 1
    for curr, totals in report.get("currencies_summary", {}).items():
        _safe_append(ws, [
            f"TOTAL {curr}",
            curr,
            totals["current"] / 100.0,
            totals["1_30"] / 100.0,
            totals["31_60"] / 100.0,
            totals["61_90"] / 100.0,
            totals["90_plus"] / 100.0,
            totals["total"] / 100.0,
        ])
        for c_idx in range(1, 9):
            c = ws.cell(row=row_num, column=c_idx)
            c.font = bold_font
            c.fill = sub_fill
            if c_idx >= 3:
                c.number_format = num_fmt
        row_num += 1

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ageing_company_block(settings: Dict[str, Any], style: Any) -> Any:
    addr = (settings.get("company_address") or "").replace("\n", "<br/>")
    phone = settings.get("company_phone") or ""
    email = settings.get("company_email") or ""
    web = settings.get("company_website") or ""
    tax_id = settings.get("tax_id") or ""
    return Paragraph(f"{addr}<br/>{phone}<br/>{email}<br/>{web}<br/>{tax_id}", style)


def export_ageing_summary_pdf(report: Dict[str, Any], settings: Dict[str, Any]) -> bytes:
    """Landscape ageing summary with letterhead, wordmark, and a full-width table."""
    buffer = io.BytesIO()
    page_w, _page_h = landscape(A4)
    left = right = 40
    usable = page_w - left - right
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=left,
        rightMargin=right,
        topMargin=36,
        bottomMargin=48,
    )
    styles = getSampleStyleSheet()
    normal = styles["Normal"]
    muted = colors.HexColor("#71717A")
    ink = colors.HexColor("#18181B")
    rule = colors.HexColor("#E4E4E7")
    hair = colors.HexColor("#A1A1AA")

    company_style = ParagraphStyle(
        "AgeingCompany",
        parent=normal,
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=muted,
        alignment=2,
    )
    title_style = ParagraphStyle(
        "AgeingTitle",
        parent=normal,
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=ink,
    )
    meta_style = ParagraphStyle(
        "AgeingMeta",
        parent=normal,
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=muted,
        alignment=2,
    )
    th_style = ParagraphStyle(
        "AgeingTh",
        parent=normal,
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
        textColor=muted,
    )
    th_right = ParagraphStyle("AgeingThRight", parent=th_style, alignment=2)
    td_style = ParagraphStyle(
        "AgeingTd",
        parent=normal,
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=ink,
    )
    td_right = ParagraphStyle("AgeingTdRight", parent=td_style, alignment=2)
    td_muted = ParagraphStyle("AgeingTdMuted", parent=td_style, textColor=muted)
    bold_style = ParagraphStyle("AgeingBold", parent=td_style, fontName="Helvetica-Bold")
    bold_right = ParagraphStyle("AgeingBoldRight", parent=td_right, fontName="Helvetica-Bold")
    empty_style = ParagraphStyle("AgeingEmpty", parent=td_style, textColor=muted, alignment=1)

    story: List[Any] = []
    as_of = report.get("as_of_date", "")
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    rows = report.get("rows") or []
    summaries = report.get("currencies_summary") or {}

    header = Table(
        [[_wordmark_flowable(normal, settings or {}), _ageing_company_block(settings or {}, company_style)]],
        colWidths=[usable * 0.55, usable * 0.45],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header)
    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.6, color=rule, spaceBefore=0, spaceAfter=14))

    title_row = Table(
        [[
            Paragraph("Receivables Ageing Summary", title_style),
            Paragraph(
                f"As of {as_of}<br/><font size='8'>Generated {generated}</font>",
                meta_style,
            ),
        ]],
        colWidths=[usable * 0.62, usable * 0.38],
    )
    title_row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(title_row)
    story.append(Spacer(1, 16))

    table_data = [[
        Paragraph("Client", th_style),
        Paragraph("Curr.", th_style),
        Paragraph("Current", th_right),
        Paragraph("1–30 days", th_right),
        Paragraph("31–60 days", th_right),
        Paragraph("61–90 days", th_right),
        Paragraph("90+ days", th_right),
        Paragraph("Total due", th_right),
    ]]

    if not rows:
        table_data.append([
            Paragraph("No outstanding receivables for this period.", empty_style),
            "", "", "", "", "", "", "",
        ])
    else:
        for r in rows:
            curr = r["currency"]
            table_data.append([
                Paragraph(r["client_name"], td_style),
                Paragraph(curr, td_muted),
                Paragraph(format_currency_amount(r["current"], curr), td_right),
                Paragraph(format_currency_amount(r["1_30"], curr), td_right),
                Paragraph(format_currency_amount(r["31_60"], curr), td_right),
                Paragraph(format_currency_amount(r["61_90"], curr), td_right),
                Paragraph(format_currency_amount(r["90_plus"], curr), td_right),
                Paragraph(format_currency_amount(r["total"], curr), bold_right),
            ])

    total_start = len(table_data)
    for curr, totals in summaries.items():
        table_data.append([
            Paragraph(f"Total {curr}", bold_style),
            Paragraph(curr, bold_style),
            Paragraph(format_currency_amount(totals["current"], curr), bold_right),
            Paragraph(format_currency_amount(totals["1_30"], curr), bold_right),
            Paragraph(format_currency_amount(totals["31_60"], curr), bold_right),
            Paragraph(format_currency_amount(totals["61_90"], curr), bold_right),
            Paragraph(format_currency_amount(totals["90_plus"], curr), bold_right),
            Paragraph(format_currency_amount(totals["total"], curr), bold_right),
        ])

    col_widths = [
        usable * 0.26,
        usable * 0.07,
        usable * 0.11,
        usable * 0.11,
        usable * 0.11,
        usable * 0.11,
        usable * 0.11,
        usable * 0.12,
    ]
    summary_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 1), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, -1), 8),
        ("LEFTPADDING", (1, 0), (-1, -1), 6),
        ("RIGHTPADDING", (1, 0), (-1, -1), 0),
        ("LINEBELOW", (0, 0), (-1, 0), 1, ink),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, rule),
    ]
    if not rows:
        style_cmds.append(("SPAN", (0, 1), (-1, 1)))
        style_cmds.append(("TOPPADDING", (0, 1), (-1, 1), 18))
        style_cmds.append(("BOTTOMPADDING", (0, 1), (-1, 1), 18))
    for i in range(1, total_start):
        if i % 2 == 0:
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")))
    if summaries:
        style_cmds.append(("LINEABOVE", (0, total_start), (-1, total_start), 0.8, hair))
        style_cmds.append(("LINEBELOW", (0, -1), (-1, -1), 1, ink))
        style_cmds.append(("BACKGROUND", (0, total_start), (-1, -1), colors.HexColor("#FAFAFA")))

    summary_table.setStyle(TableStyle(style_cmds))
    story.append(summary_table)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setStrokeColor(rule)
        canvas.setLineWidth(0.5)
        canvas.line(left, 32, page_w - right, 32)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        canvas.drawString(left, 20, f"{(settings or {}).get('company_name') or ''}  ·  Confidential".strip(" ·"))
        canvas.restoreState()

    doc.build(story, canvasmaker=NumberedCanvas, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def _period_label(date_from: str, date_to: str) -> str:
    return f"Period {date_from} to {date_to}"


def _standard_report_pdf(
    *,
    title: str,
    period_label: str,
    settings: Dict[str, Any],
    headers: List[str],
    col_weights: List[float],
    rows: List[List[str]],
    right_align_from: int = 1,
    kpi_items: Optional[List[tuple]] = None,
    landscape_page: bool = True,
    extra_flowables: Optional[List[Any]] = None,
) -> bytes:
    """Letterheaded tabular report with wordmark, period, and confidential footer."""
    finance_pdf_mod._ensure_pdf_fonts()
    font_regular = finance_pdf_mod.FONT_REGULAR
    font_bold = finance_pdf_mod.FONT_BOLD
    buffer = io.BytesIO()
    pagesize = landscape(A4) if landscape_page else A4
    page_w, _page_h = pagesize
    left = right = 40
    usable = page_w - left - right
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=left,
        rightMargin=right,
        topMargin=36,
        bottomMargin=48,
    )
    styles = getSampleStyleSheet()
    normal = styles["Normal"]
    muted = colors.HexColor("#71717A")
    ink = colors.HexColor("#18181B")
    rule = colors.HexColor("#E4E4E7")
    hair = colors.HexColor("#A1A1AA")

    company_style = ParagraphStyle(
        "StdCompany", parent=normal, fontName=font_regular, fontSize=8, leading=11, textColor=muted, alignment=2,
    )
    title_style = ParagraphStyle(
        "StdTitle", parent=normal, fontName=font_bold, fontSize=16, leading=20, textColor=ink,
    )
    meta_style = ParagraphStyle(
        "StdMeta", parent=normal, fontName=font_regular, fontSize=8.5, leading=12, textColor=muted, alignment=2,
    )
    th_style = ParagraphStyle(
        "StdTh", parent=normal, fontName=font_regular, fontSize=7.5, leading=10, textColor=muted,
    )
    th_right = ParagraphStyle("StdThRight", parent=th_style, alignment=2)
    td_style = ParagraphStyle(
        "StdTd", parent=normal, fontName=font_regular, fontSize=8.5, leading=11, textColor=ink,
    )
    td_right = ParagraphStyle("StdTdRight", parent=td_style, alignment=2)
    kpi_label = ParagraphStyle(
        "StdKpiLabel", parent=normal, fontName=font_regular, fontSize=7.5, leading=10, textColor=muted,
    )
    empty_style = ParagraphStyle("StdEmpty", parent=td_style, textColor=muted, alignment=1)

    story: List[Any] = []
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    header = Table(
        [[_wordmark_flowable(normal, settings or {}), _ageing_company_block(settings or {}, company_style)]],
        colWidths=[usable * 0.55, usable * 0.45],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header)
    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.6, color=rule, spaceBefore=0, spaceAfter=14))
    title_row = Table(
        [[
            Paragraph(title, title_style),
            Paragraph(f"{period_label}<br/><font size='8'>Generated {generated}</font>", meta_style),
        ]],
        colWidths=[usable * 0.62, usable * 0.38],
    )
    title_row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(title_row)
    story.append(Spacer(1, 14))

    if kpi_items:
        kpi_table_data = [[
            Paragraph(f"{label}<br/><font color='#18181B' size='11'><b>{value}</b></font>", kpi_label)
            for label, value in kpi_items
        ]]
        kpi_w = usable / max(len(kpi_items), 1)
        kpi_table = Table(kpi_table_data, colWidths=[kpi_w] * len(kpi_items))
        kpi_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
            ("BOX", (0, 0), (-1, -1), 0.4, rule),
            ("LINEAFTER", (0, 0), (-2, -1), 0.4, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 14))

    if extra_flowables:
        story.extend(extra_flowables)

    header_row = []
    for idx, label in enumerate(headers):
        style = th_right if idx >= right_align_from else th_style
        header_row.append(Paragraph(label, style))
    table_data: List[List[Any]] = [header_row]
    if not rows:
        empty = [Paragraph("No activity in this period.", empty_style)] + [""] * (len(headers) - 1)
        table_data.append(empty)
    else:
        for row in rows:
            cells = []
            for idx, cell in enumerate(row):
                style = td_right if idx >= right_align_from else td_style
                cells.append(Paragraph(str(cell), style))
            table_data.append(cells)

    weight_sum = sum(col_weights) or 1
    col_widths = [usable * (w / weight_sum) for w in col_weights]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, -1), 8),
        ("LEFTPADDING", (1, 0), (-1, -1), 6),
        ("RIGHTPADDING", (1, 0), (-1, -1), 0),
        ("LINEBELOW", (0, 0), (-1, 0), 1, ink),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, rule),
    ]
    if not rows:
        style_cmds.append(("SPAN", (0, 1), (-1, 1)))
        style_cmds.append(("TOPPADDING", (0, 1), (-1, 1), 18))
        style_cmds.append(("BOTTOMPADDING", (0, 1), (-1, 1), 18))
    else:
        for i in range(1, len(table_data)):
            if i % 2 == 0:
                style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")))
        style_cmds.append(("LINEBELOW", (0, -1), (-1, -1), 1, hair))
    table.setStyle(TableStyle(style_cmds))
    story.append(table)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setStrokeColor(rule)
        canvas.setLineWidth(0.5)
        canvas.line(left, 32, page_w - right, 32)
        canvas.setFont(font_regular, 7)
        canvas.setFillColor(muted)
        canvas.drawString(left, 20, f"{(settings or {}).get('company_name') or ''}  ·  Confidential".strip(" ·"))
        canvas.restoreState()

    doc.build(story, canvasmaker=NumberedCanvas, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def export_receivables_outstanding_pdf(report: Dict[str, Any], settings: Dict[str, Any]) -> bytes:
    rows = [
        [r.get("client_name") or "", r.get("currency") or "PKR", str(r.get("invoice_count") or 0),
         format_currency_amount(_as_int(r.get("total")), r.get("currency") or "PKR")]
        for r in report.get("rows") or []
    ]
    for curr, total in (report.get("totals") or {}).items():
        rows.append([f"Total {curr}", curr, "", format_currency_amount(_as_int(total), curr)])
    kpis = [
        (f"Outstanding {curr}", format_currency_amount(_as_int(total), curr))
        for curr, total in (report.get("totals") or {}).items()
    ]
    return _standard_report_pdf(
        title="Receivables Outstanding",
        period_label=_period_label(report.get("date_from") or "", report.get("date_to") or ""),
        settings=settings,
        headers=["Client", "Currency", "Open invoices", "Amount due"],
        col_weights=[0.46, 0.14, 0.16, 0.24],
        rows=rows,
        right_align_from=2,
        kpi_items=kpis or None,
        landscape_page=False,
    )


def export_collections_pdf(report: Dict[str, Any], settings: Dict[str, Any]) -> bytes:
    summary = report.get("summary") or {}
    rows = []
    for r in report.get("receipts") or []:
        method = "Cheque" if r.get("is_cheque") else (r.get("method") or "").capitalize()
        if r.get("bank"):
            method = f"{method} ({str(r.get('bank')).upper()})"
        status = "Approved" if r.get("status") == "posted" else (r.get("status") or "").capitalize()
        rows.append([
            r.get("date") or "",
            r.get("number") or "",
            r.get("client_name") or "",
            method,
            status,
            format_currency_amount(_as_int(r.get("amount_received")), r.get("currency") or "PKR"),
            format_currency_amount(_as_int(r.get("pkr_amount")), "PKR"),
        ])
    kpis = [
        ("Total collection", format_currency_amount(_as_int(summary.get("total_collected_pkr")), "PKR")),
        ("Pending cheques", format_currency_amount(_as_int(summary.get("pending_pkr")), "PKR")),
        ("Cash", format_currency_amount(_as_int(summary.get("cash_pkr")), "PKR")),
        ("Bank (ABL + Meezan)", format_currency_amount(
            _as_int(summary.get("abl_pkr")) + _as_int(summary.get("meezan_pkr")), "PKR"
        )),
    ]
    return _standard_report_pdf(
        title="Collections Report",
        period_label=_period_label(report.get("date_from") or "", report.get("date_to") or ""),
        settings=settings,
        headers=["Date", "Receipt", "Client", "Method", "Status", "Amount", "PKR equivalent"],
        col_weights=[0.11, 0.13, 0.22, 0.16, 0.10, 0.14, 0.14],
        rows=rows,
        right_align_from=5,
        kpi_items=kpis,
        landscape_page=True,
    )


def export_summary_accounts_pdf(report: Dict[str, Any], settings: Dict[str, Any]) -> bytes:
    rows = []
    for idx, r in enumerate(report.get("rows") or [], start=1):
        curr = r.get("currency") or "PKR"
        rows.append([
            str(idx),
            r.get("issue_date") or "",
            r.get("number") or "",
            r.get("client_name") or "",
            curr,
            format_currency_amount(_as_int(r.get("invoice_amount")), curr),
            format_currency_amount(_as_int(r.get("received")), curr),
            format_currency_amount(_as_int(r.get("balance_due")), curr),
        ])
    for curr, totals in (report.get("totals") or {}).items():
        rows.append([
            "",
            "",
            "",
            f"Total {curr}",
            curr,
            format_currency_amount(_as_int(totals.get("invoice")), curr),
            format_currency_amount(_as_int(totals.get("received")), curr),
            format_currency_amount(_as_int(totals.get("due")), curr),
        ])
    kpis = []
    for curr, totals in (report.get("totals") or {}).items():
        kpis.append((f"Invoiced {curr}", format_currency_amount(_as_int(totals.get("invoice")), curr)))
        kpis.append((f"Received {curr}", format_currency_amount(_as_int(totals.get("received")), curr)))
        kpis.append((f"Due {curr}", format_currency_amount(_as_int(totals.get("due")), curr)))
    return _standard_report_pdf(
        title="Summary Accounts",
        period_label=_period_label(report.get("date_from") or "", report.get("date_to") or ""),
        settings=settings,
        headers=["Sr", "Date", "Invoice", "Client", "Curr.", "Invoice amount", "Amount received", "Balance due"],
        col_weights=[0.06, 0.11, 0.12, 0.22, 0.07, 0.14, 0.14, 0.14],
        rows=rows,
        right_align_from=5,
        kpi_items=kpis[:6] or None,
        landscape_page=True,
    )
