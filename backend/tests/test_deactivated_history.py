"""Deactivated members keep their history visible (multi-company mode).

Throwaway Mongo only (name contains 'test'); outbound network blocked.
"""
from __future__ import annotations

import os
import socket
import sys
import unittest
from unittest import mock

os.environ.setdefault("SECRET_KEY", "test-secret-key-deactivated-history-0123456789")
os.environ["RUN_SCHEDULERS"] = "0"
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import httpx

from app.config import settings
from app.core.security import claims_with_session, create_access_token
from app.database import connect_to_mongo
from app.main import app

TEST_DB = "reamarc_test_deactivated_history"
CID = "company_test_dh"
DEACT = "2026-10-15"

_real_connect = socket.socket.connect


def _loopback_only(self, addr, *a, **k):
    host = addr[0] if isinstance(addr, tuple) and addr else addr
    if host in ("127.0.0.1", "::1", "localhost"):
        return _real_connect(self, addr, *a, **k)
    raise OSError("outbound blocked in tests")


ROLES = {"dh_adm": "admin", "dh_hikmat": "team_member", "dh_emp": "team_member", "dh_gone": "team_member"}
HIKMAT_DAYS = ["2026-10-01", "2026-10-02", "2026-10-06", "2026-10-07", "2026-10-08", "2026-10-13", "2026-10-14"]


class TestDeactivatedHistory(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        for p in (
            mock.patch.object(socket.socket, "connect", new=_loopback_only),
            mock.patch.object(settings, "MULTI_COMPANY_ENABLED", True),
            mock.patch("app.services.push_service.send_expo_push", new=mock.AsyncMock(return_value=0)),
            mock.patch("app.services.web_push_service.send_web_push", new=mock.AsyncMock(return_value=0)),
        ):
            p.start()
            self.addCleanup(p.stop)
        object.__setattr__(settings, "MONGODB_DB_NAME", TEST_DB)
        assert "test" in TEST_DB
        await connect_to_mongo()
        from app.database import db_instance

        self.db = db_instance.db
        await db_instance.client.drop_database(TEST_DB)
        from app.core.limiter import limiter

        limiter.enabled = False
        self.company = {"id": CID, "name": CID, "slug": CID, "status": "active", "token_version": 0}
        await self.db.companies.insert_one(dict(self.company))
        self.users = {}
        for uid, role in ROLES.items():
            u = {"id": uid, "email": f"{uid}@reamarc-test.com", "full_name": uid, "name": uid, "role": role,
                 "is_active": True, "company_id": CID, "session_epoch": 0, "workspace_ids": [],
                 "department": "Content", "departments": ["Content"], "primary_department": "Content",
                 "joining_date": "2026-09-01", "refresh_jtis": [f"web-{uid}"], "refresh_retired": []}
            m = {"id": f"m_{uid}", "user_id": uid, "company_id": CID, "status": "active", "role": role,
                 "workspace_ids": [], "token_version": 0, "department": "Content",
                 "primary_department": "Content", "departments": ["Content"]}
            await self.db.users.insert_one(dict(u))
            await self.db.company_memberships.insert_one(dict(m))
            self.users[uid] = (u, m)
        for d in HIKMAT_DAYS:
            await self.db.attendance_records.insert_one({
                "id": f"att_{d}", "user_id": "dh_hikmat", "date": d, "status": "present",
                "check_in": f"{d}T11:00:00+05:00", "check_out": f"{d}T19:00:00+05:00",
                "work_hours": 8.0, "working_hours_minutes": 480, "company_id": CID,
            })
        await self.db.daily_log_entries.insert_one({"id": "log1", "user_id": "dh_hikmat", "date": "2026-10-14",
                                                    "resource_name": "dh_hikmat", "company_id": CID})
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")
        # Deactivate Hikmat the way Team directory does it.
        r = await self.client.patch("/api/v1/admin/members/dh_hikmat",
                                    json={"is_active": False, "deactivation_date": DEACT},
                                    headers=self._auth("dh_adm"))
        self.assertEqual(r.status_code, 200, r.text)

    async def asyncTearDown(self):
        await self.client.aclose()
        from app.database import db_instance

        await db_instance.client.drop_database(TEST_DB)
        db_instance.client = None
        db_instance.db = None

    def _auth(self, uid):
        u, m = self.users[uid]
        base = {"sub": uid, "email": u["email"], "name": u["full_name"], "role": u["role"], "cid": CID,
                "workspace_ids": [], "mid": m["id"], "mtv": 0}
        c = claims_with_session(u, base, company_doc=self.company, membership_doc=m)
        return {"Authorization": f"Bearer {create_access_token(dict(c, sid=f'web-{uid}'))}",
                "X-Requested-With": "XMLHttpRequest"}

    async def _in_company(self, coro_fn):
        from app.core.tenant import TenantContext, set_tenant_context, reset_tenant_context
        tok = set_tenant_context(TenantContext(company_id=CID))
        try:
            return await coro_fn()
        finally:
            reset_tenant_context(tok)

    async def test_deactivation_keeps_data_and_records_the_date(self):
        mem = await self.db.company_memberships.find_one({"user_id": "dh_hikmat", "company_id": CID})
        user = await self.db.users.find_one({"id": "dh_hikmat"})
        self.assertEqual(mem["status"], "inactive")
        self.assertEqual(mem["deactivation_date"], DEACT)
        self.assertIs(user["is_active"], False)
        self.assertEqual(await self.db.attendance_records.count_documents({"user_id": "dh_hikmat"}), len(HIKMAT_DAYS))
        self.assertEqual(await self.db.daily_log_entries.count_documents({"user_id": "dh_hikmat"}), 1)

    async def test_team_directory_lists_him_as_inactive(self):
        r = await self.client.get("/api/v1/admin/members", headers=self._auth("dh_adm"))
        self.assertEqual(r.status_code, 200, r.text)
        by_id = {m["id"]: m for m in r.json()}
        self.assertIn("dh_hikmat", by_id)
        self.assertIs(by_id["dh_hikmat"]["is_active"], False)
        self.assertEqual(by_id["dh_hikmat"].get("deactivation_date"), DEACT)
        r = await self.client.get("/api/v1/admin/members?is_active=true", headers=self._auth("dh_adm"))
        self.assertNotIn("dh_hikmat", {m["id"] for m in r.json()})
        r = await self.client.get("/api/v1/admin/members?is_active=false", headers=self._auth("dh_adm"))
        self.assertEqual([m["id"] for m in r.json()], ["dh_hikmat"])

    async def test_punctuality_keeps_him_for_months_he_worked(self):
        from app.services import attendance_service as svc
        oct_ = await self._in_company(lambda: svc.get_monthly_punctuality_summary(2026, 10))
        rows = {r.user_id: r for r in oct_.rows}
        self.assertIn("dh_hikmat", rows)
        self.assertEqual(rows["dh_hikmat"].days_present, len(HIKMAT_DAYS))
        # Working days stop at the deactivation date (Oct 1-14: 11 working days, first Saturday off).
        full = rows["dh_emp"].total_working_days
        self.assertLess(rows["dh_hikmat"].total_working_days, full)
        self.assertEqual(rows["dh_hikmat"].total_working_days, 11)
        nov = await self._in_company(lambda: svc.get_monthly_punctuality_summary(2026, 11))
        self.assertNotIn("dh_hikmat", {r.user_id for r in nov.rows})

    async def test_daily_register_before_and_after_deactivation(self):
        from app.services import attendance_service as svc
        before = await self._in_company(lambda: svc.get_daily_matrix("2026-10-14"))
        self.assertIn("dh_hikmat", {r.user_id for r in before.rows})
        after = await self._in_company(lambda: svc.get_daily_matrix("2026-10-16"))
        self.assertNotIn("dh_hikmat", {r.user_id for r in after.rows})

    async def test_his_timesheet_still_opens(self):
        from app.services import attendance_service as svc
        ts = await self._in_company(lambda: svc.get_timesheet_for_user_id("dh_hikmat", 2026, 10))
        dated = {r.date for r in ts.records if getattr(r, "check_in", None)}
        self.assertTrue(set(HIKMAT_DAYS) <= dated, dated)

    async def test_excel_export_has_his_tab(self):
        from app.services.attendance_excel import generate_multi_tab_attendance_workbook
        from openpyxl import load_workbook
        buf = await self._in_company(lambda: generate_multi_tab_attendance_workbook(2026, 10))
        wb = load_workbook(buf)
        self.assertTrue(any("dh_hikmat" in n for n in wb.sheetnames), wb.sheetnames)

    async def test_deleting_someone_else_does_not_wipe_his_history(self):
        r = await self.client.delete("/api/v1/admin/members/dh_gone", headers=self._auth("dh_adm"))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(await self.db.attendance_records.count_documents({"user_id": "dh_hikmat"}), len(HIKMAT_DAYS))
        self.assertEqual(await self.db.daily_log_entries.count_documents({"user_id": "dh_hikmat"}), 1)

    async def test_new_work_lists_still_exclude_him(self):
        from app.services.company_members import company_member_users, company_member_user
        from app.services.access import tracked_users

        async def go():
            default = await company_member_users({})
            active = await company_member_users({"is_active": {"$ne": False}})
            tracked = await tracked_users(None or __import__("app.database", fromlist=["x"]).get_database(), "attendance")
            one = await company_member_user("dh_hikmat")
            return default, active, tracked, one
        default, active, tracked, one = await self._in_company(go)
        for lst in (default, active, tracked):
            self.assertNotIn("dh_hikmat", {u.get("id") for u in lst})
        self.assertIsNone(one)

    async def test_reactivating_restores_him(self):
        r = await self.client.patch("/api/v1/admin/members/dh_hikmat", json={"is_active": True},
                                    headers=self._auth("dh_adm"))
        self.assertEqual(r.status_code, 200, r.text)
        r = await self.client.get("/api/v1/admin/members?is_active=true", headers=self._auth("dh_adm"))
        self.assertIn("dh_hikmat", {m["id"] for m in r.json()})
        from app.services import attendance_service as svc
        nov = await self._in_company(lambda: svc.get_daily_matrix("2026-10-16"))
        self.assertIn("dh_hikmat", {r.user_id for r in nov.rows})


if __name__ == "__main__":
    unittest.main()
